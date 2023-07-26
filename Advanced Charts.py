from openpyxl import Workbook
from openpyxl.chart import BarChart, Series, Reference

wb = Workbook(write_only=True)
ws = wb.create_sheet('transactions.xlsx')

rows = [
    ('Number', '2022', '2023'),
    (2, 100000, 300000),
    (3, 450000, 750000),
    (4, 145000, 340000),
    (5, 200000, 123400),
    (6, 238400, 401408),
    (7, 530590, 300387),
]


for row in rows:
    ws.append(row)


chart1 = BarChart()
chart1.type = "col"
chart1.style = 10
chart1.title = "Bar Chart"
chart1.y_axis.title = 'Initial'
chart1.x_axis.title = 'Growth'

data = Reference(ws, min_col=2, min_row=1, max_row=7, max_col=3)
cats = Reference(ws, min_col=1, min_row=2, max_row=7)
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(cats)
chart1.shape = 4
ws.add_chart(chart1, "A10")
wb.save('transactions.xlsx')